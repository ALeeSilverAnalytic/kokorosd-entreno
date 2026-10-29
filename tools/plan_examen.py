#!/usr/bin/env python3
"""Plan de examen: plantilla Excel <-> content/examen_1dan.json
  python3 tools/plan_examen.py plantilla  salida.xlsx      (genera la plantilla con el plan propuesto)
  python3 tools/plan_examen.py importar   entrada.xlsx     (lee el Excel editado y escribe content/examen_1dan.json)
Columnas del Excel: Fase | Tarea | Inicio | Fin | Quién | Notas   (fechas dd/mm/aaaa; Quién: TODOS o alias separados por coma)"""
import sys, json, uuid, datetime as dt, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
NS = uuid.UUID('6f1c1d0e-4b7a-4c55-9d52-1c0a0c0c0c0c')
EXAMEN = {'slug': '1dan-2027', 'nombre': 'Examen 1er DAN', 'fecha': '2027-06-15'}
D = lambda s: dt.date.fromisoformat(s)
PLAN = [  # fase, tarea, inicio, fin, quién, notas
 ('0 · Lanzamiento', 'Reunión de lanzamiento', '2026-11-02', '2026-11-02', 'TODOS', ''),
 ('0 · Lanzamiento', 'Reparto del archivo Gantt', '2026-11-02', '2026-11-04', 'TODOS', ''),
 ('0 · Lanzamiento', 'Envío de los PDF del programa', '2026-11-09', '2026-11-09', 'TODOS', ''),
 ('0 · Lanzamiento', 'Detectar y marcar lo que no está en vídeo', '2026-11-09', '2026-11-22', 'TODOS', ''),
 ('0 · Lanzamiento', 'Videoconferencia de coordinación', '2026-11-23', '2026-11-23', 'TODOS', ''),
 ('1 · Repaso de cinturones inferiores', 'Repaso cinturón amarillo', '2026-11-23', '2026-12-06', 'TODOS', ''),
 ('1 · Repaso de cinturones inferiores', 'Repaso cinturón naranja', '2026-12-07', '2026-12-20', 'TODOS', ''),
 ('1 · Repaso de cinturones inferiores', 'Repaso cinturón verde', '2026-12-21', '2027-01-03', 'TODOS', ''),
 ('1 · Repaso de cinturones inferiores', 'Repaso cinturón azul', '2027-01-04', '2027-01-17', 'TODOS', ''),
 ('1 · Repaso de cinturones inferiores', 'Repaso cinturón marrón', '2027-01-18', '2027-01-31', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Material del profesor', '2027-02-01', '2027-02-07', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Agresiones especiales', '2027-02-01', '2027-02-21', 'TODOS', 'Sentado en silla, de pie contra pared y tendido en el suelo'),
 ('2 · Programa 1er DAN', 'Posturas y guardias básicas', '2027-02-22', '2027-02-28', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Atemis marchando', '2027-03-01', '2027-03-07', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Puntos vitales', '2027-03-08', '2027-03-28', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Combate total', '2027-03-29', '2027-04-11', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Manejo del yawara', '2027-04-12', '2027-05-09', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Agilidad', '2027-05-10', '2027-05-16', 'TODOS', ''),
 ('2 · Programa 1er DAN', 'Teoría', '2027-05-17', '2027-05-30', 'TODOS', ''),
 ('3 · Trabajo continuo', 'Katas (todos los de los cinturones inferiores y los del 1er DAN)', '2027-02-01', '2027-06-08', 'TODOS', ''),
 ('3 · Trabajo continuo', 'Endurecimiento para rompimientos', '2027-02-01', '2027-06-08', 'TODOS', ''),
 ('3 · Trabajo continuo', 'Comprar tablas', '2027-04-26', '2027-05-09', 'TODOS', ''),
 ('4 · Cierre', 'Pago de tasas', '2027-04-01', '2027-04-30', 'TODOS', 'Fecha provisional'),
 ('4 · Cierre', 'Concentración presencial', '2027-05-22', '2027-05-23', 'TODOS', 'Fecha provisional'),
 ('4 · Cierre', 'Repaso general', '2027-05-31', '2027-06-14', 'TODOS', ''),
 ('4 · Cierre', 'EXAMEN 1er DAN', '2027-06-15', '2027-06-15', 'TODOS', ''),
]
def plantilla(out):
    wb = Workbook(); ws = wb.active; ws.title = 'Plan'
    ws.append(['Fase', 'Tarea', 'Inicio', 'Fin', 'Quién', 'Notas'])
    for f in ws[1]: f.font = Font(bold=True, color='FFFFFF'); f.fill = PatternFill('solid', fgColor='2F5D50')
    for fase, t, i, f, q, n in PLAN: ws.append([fase, t, D(i), D(f), q, n])
    for r in ws.iter_rows(min_row=2):
        r[2].number_format = r[3].number_format = 'dd/mm/yyyy'
    for c, w in zip('ABCDEF', (34, 62, 13, 13, 16, 52)): ws.column_dimensions[c].width = w
    ws.freeze_panes = 'A2'
    h = wb.create_sheet('Cómo editar')
    for l in ['Examen: ' + EXAMEN['nombre'] + ' · fecha ' + D(EXAMEN['fecha']).strftime('%d/%m/%Y'),
              'Edita la hoja "Plan": cambia fechas, borra o añade filas, cambia el orden.',
              'Fase: agrupa las tareas en la app (usa el mismo texto en las tareas de la misma fase).',
              'Quién: TODOS, o alias separados por coma (por ejemplo: shark, sr.smith). Solo para tu control; los que se examinan ven el plan completo.',
              'Fechas en formato dd/mm/aaaa. Para un solo día, Inicio y Fin iguales.',
              'Cuando termines, envíame el archivo y lo cargo en la app.']: h.append([l])
    h.column_dimensions['A'].width = 120
    wb.save(out)
def importar(path):
    ws = load_workbook(path, data_only=True)['Plan']; tareas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        fase, t, i, f, q, n = (list(r) + [None] * 6)[:6]
        if not t: continue
        g = lambda v: v.date() if isinstance(v, dt.datetime) else v if isinstance(v, dt.date) else dt.datetime.strptime(str(v).strip(), '%d/%m/%Y').date()
        tareas.append({'orden': len(tareas) + 1, 'fase': str(fase or '').strip(), 'titulo': str(t).strip(), 'inicio': g(i).isoformat(), 'fin': g(f or i).isoformat(), 'quien': str(q or 'TODOS').strip(), 'nota': (str(n).strip() if n else None)})
    escribir(tareas)
def escribir(tareas):
    eid = str(uuid.uuid5(NS, 'examen-' + EXAMEN['slug']))
    rows = [{'id': str(uuid.uuid5(NS, f"examen-tarea-{EXAMEN['slug']}-{t['orden']}")), 'examen_slug': EXAMEN['slug'], **t} for t in tareas]
    ops = [{'t': 'examen', 'k': 'slug', 'rows': [{'id': eid, **EXAMEN}]}, {'t': 'examen_tarea', 'k': 'id', 'rows': rows}]
    here = os.path.dirname(os.path.abspath(__file__))
    json.dump({'ops': ops}, open(os.path.join(here, '..', 'content', 'examen_1dan.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
    print(len(rows), 'tareas escritas')
if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == 'plantilla': plantilla(sys.argv[2])
    elif len(sys.argv) == 3 and sys.argv[1] == 'importar': importar(sys.argv[2])
    elif len(sys.argv) == 2 and sys.argv[1] == 'propuesta': escribir([{'orden': i + 1, 'fase': a, 'titulo': b, 'inicio': c, 'fin': d, 'quien': e, 'nota': f or None} for i, (a, b, c, d, e, f) in enumerate(PLAN)])
    else: print(__doc__)
